#!/usr/bin/env python3
"""Le stringhe dell'interfaccia dal file dei partner, dentro la pagina.

`StratiGraph_UI_strings.xlsx` (OneDrive, WP 01 Resources, «UI translations») è
il file in cui i partner traducono. Questo script lo LEGGE e non lo scrive mai:
prende le colonne delle lingue chieste e riscrive, in `web/index.html`, il blocco
fra i due segni `>>> ui_strings.py` e `<<< ui_strings.py` dentro `STRINGS`.

Le stringhe restano INLINE nella pagina, com'è per progetto: un file HTML, nessun
passo di build sul telefono, niente da scaricare a tre metri sotto terra. Lo
script è il passo di build che sta sulla scrivania.

CHE COSA ENTRA, e perché così poco:

* solo le chiavi che la pagina ha (`STRINGS.en`) — l'xlsx copre anche superfici
  che qui non ci sono;
* solo dove l'INGLESE dell'xlsx è lo stesso della pagina: una stessa chiave con
  due inglesi diversi è una traduzione di un'altra frase («StratiField · desk»
  non è «field assistant»). Si salta e si dice;
* solo se i `{segnaposto}` e i termini del dominio (US, ORCID, …) sopravvivono:
  una frase che perde `{n}` dice «nota di» e si ferma.

Il resto ricade sull'inglese, come per ogni chiave senza traduzione.

    python scripts/ui_strings.py                 # he, de → web/index.html
    python scripts/ui_strings.py --langs he de ro
    python scripts/ui_strings.py --check         # esce 1 se la pagina è indietro
    python scripts/ui_strings.py --report        # chiavi che mancano all'xlsx
"""

from __future__ import annotations

import argparse
import json
import os
import pathlib
import re
import sys
from typing import Dict, List, Tuple

ROOT = pathlib.Path(__file__).resolve().parent.parent
PAGE = ROOT / "web" / "index.html"

DEFAULT_XLSX = pathlib.Path(os.environ.get("SG_UI_STRINGS_XLSX") or (
    pathlib.Path.home() / "Library/CloudStorage/OneDrive-CNR/0_RACCOLTA_CNR"
    / "StratiGraph-grp - Documenti"
    / "WP 01 - Management, communication, dissemination"
    / "Resources (visual identity guidebook, templates, logo, etc)"
    / "UI translations" / "StratiGraph_UI_strings.xlsx"))

SHEET = "UI strings"
DEFAULT_LANGS = ("he", "de")
BEGIN = "  // >>> ui_strings.py"
END = "  // <<< ui_strings.py"
#: gli stessi di `tests/test_locales.py::test_no_domain_term_was_translated`
TERMS = ("US", "DTC", "ORCID", "HDT", "em.json", "crmdig")


def page_strings(page: str) -> Dict[str, Dict[str, str]]:
    """`{locale: {key: value}}` come la pagina li dichiara (lo stesso parser dei
    test, `tests/test_field_signature.py::locales`)."""
    block = re.search(r"const STRINGS = \{(.*?)\n\};", page, re.S)
    if not block:
        raise SystemExit("la pagina non dichiara STRINGS")
    found: Dict[str, Dict[str, str]] = {}
    for match in re.finditer(r"^  (\w+): \{(.*?)^  \}", block.group(1),
                             re.S | re.M):
        found[match.group(1)] = {
            k: json.loads(f'"{v}"') for k, v in
            re.findall(r'"([^"]+)":\s*"((?:[^"\\]|\\.)*)"', match.group(2))}
    return found


def read_xlsx(path: pathlib.Path) -> Tuple[List[str], Dict[str, Dict[str, str]]]:
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(book[SHEET].iter_rows(values_only=True))
    header = [str(c or "").strip() for c in rows[0]]
    table: Dict[str, Dict[str, str]] = {}
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        cells = dict(zip(header, row))
        table[str(row[0]).strip()] = {
            k: ("" if v is None else str(v)) for k, v in cells.items()}
    return header, table


def _placeholders(text: str) -> set:
    return set(re.findall(r"\{(\w+)\}", text))


def select(en: Dict[str, str], table: Dict[str, Dict[str, str]],
           lang: str) -> Tuple[Dict[str, str], List[Tuple[str, str]]]:
    """Le stringhe di `lang` che la pagina può prendere, e quelle saltate col
    motivo."""
    taken: Dict[str, str] = {}
    skipped: List[Tuple[str, str]] = []
    for key in en:                                   # l'ordine della pagina
        row = table.get(key)
        if row is None:
            continue
        value = row.get(lang, "").strip()
        if not value:
            skipped.append((key, "vuota nell'xlsx"))
            continue
        source = row.get("en (source)", "")
        if source != en[key]:
            skipped.append((key, f"inglese diverso: pagina «{en[key]}», "
                                 f"xlsx «{source}»"))
            continue
        if _placeholders(value) != _placeholders(en[key]):
            skipped.append((key, f"segnaposto persi: «{value}»"))
            continue
        lost = [t for t in TERMS if t in en[key] and t not in value]
        if lost:
            skipped.append((key, f"termine tradotto ({', '.join(lost)}): «{value}»"))
            continue
        taken[key] = value
    return taken, skipped


def render(blocks: Dict[str, Dict[str, str]], source: pathlib.Path) -> str:
    lines = [BEGIN,
             f"  // GENERATO da scripts/ui_strings.py leggendo {source.name}",
             "  // (foglio «UI strings»): non si modifica a mano, si rilancia lo",
             "  // script. BOZZE dei partner, solo le chiavi che la pagina ha con",
             "  // lo stesso inglese; il resto ricade sull'inglese."]
    for lang, strings in blocks.items():
        lines.append(f"  {lang}: {{")
        for key, value in strings.items():
            lines.append(f"    {json.dumps(key)}: "
                         f"{json.dumps(value, ensure_ascii=False)},")
        lines.append("  },")
    lines.append(END)
    return "\n".join(lines)


def spliced(page: str, generated: str) -> str:
    if BEGIN in page:
        start = page.index(BEGIN)
        stop = page.index(END, start) + len(END)
        return page[:start] + generated + page[stop:]
    # la prima volta: PRIMA della riga degli slot vuoti (`ro: {}, el: {}, …`).
    # Dopo, no: il parser dei test legge `^  ro: {` fino al primo `^  }` che
    # trova, e con un blocco sotto si sarebbe mangiato l'ebraico (misurato).
    slots = re.search(r"^  // ── the partners' slots.*$", page, re.M)
    if slots:
        return page[:slots.start()] + generated + "\n" + page[slots.start():]
    block = re.search(r"const STRINGS = \{.*?\n\};", page, re.S)
    close = block.end() - len("};")
    return page[:close] + generated + "\n" + page[close:]


def main(argv: List[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--xlsx", type=pathlib.Path, default=DEFAULT_XLSX)
    parser.add_argument("--langs", nargs="+", default=list(DEFAULT_LANGS))
    parser.add_argument("--check", action="store_true",
                        help="non scrive: esce 1 se la pagina non è allineata")
    parser.add_argument("--report", action="store_true",
                        help="elenca le chiavi della pagina che l'xlsx non ha")
    args = parser.parse_args(argv)

    if not args.xlsx.is_file():
        print(f"xlsx non trovato: {args.xlsx}", file=sys.stderr)
        return 2
    page = PAGE.read_text(encoding="utf-8")
    en = page_strings(page)["en"]
    header, table = read_xlsx(args.xlsx)
    missing_cols = [lang for lang in args.langs if lang not in header]
    if missing_cols:
        print(f"l'xlsx non ha le colonne {missing_cols}", file=sys.stderr)
        return 2

    blocks: Dict[str, Dict[str, str]] = {}
    for lang in args.langs:
        taken, skipped = select(en, table, lang)
        blocks[lang] = taken
        print(f"{lang}: {len(taken)}/{len(en)} chiavi dall'xlsx, "
              f"{len(skipped)} saltate")
        for key, why in skipped:
            print(f"   salto {key}: {why}")

    if args.report:
        absent = [k for k in en if k not in table]
        print(f"\n{len(absent)} chiavi della pagina che l'xlsx NON ha:")
        for key in absent:
            print(f"   {key}\t{en[key]}")

    updated = spliced(page, render(blocks, args.xlsx))
    if args.check:
        if updated != page:
            print("web/index.html NON è allineata all'xlsx: rilancia lo script",
                  file=sys.stderr)
            return 1
        print("web/index.html allineata all'xlsx")
        return 0
    if updated != page:
        PAGE.write_text(updated, encoding="utf-8")
        print("web/index.html aggiornata")
    else:
        print("web/index.html già allineata")
    return 0


if __name__ == "__main__":
    sys.exit(main())
